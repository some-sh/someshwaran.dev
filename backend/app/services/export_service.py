"""Renders a Resume into downloadable PDF/DOCX bytes.

Per CLAUDE.md's architecture principle, this is the ONLY place resume
content becomes a document: both the public default-resume export and
the admin per-resume export (app/api/public.py, app/api/admin.py) call
into these two functions, so there is exactly one rendering path per
output format for the whole app — never one copy of the layout per
route.
"""

from __future__ import annotations

import io
import re

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.models.resume import Experience, Resume


def slugify_filename(name: str) -> str:
    """'Backend-focused (copy)' -> 'backend-focused-copy', used as the
    download filename stem for both export formats."""
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return slug or "resume"


def _contact_line(resume: Resume) -> str:
    info = resume.personal_info
    if info is None:
        return ""
    fields = (info.email, info.phone, info.location, info.website, info.github, info.linkedin)
    parts = [p for p in fields if p]
    return " • ".join(parts)


def _experience_dates(exp: Experience) -> str:
    end = exp.end_date.isoformat() if exp.end_date else "Present"
    return f"{exp.start_date.isoformat()} – {end}"


def _skills_by_category(resume: Resume) -> dict[str, list[str]]:
    by_category: dict[str, list[str]] = {}
    for skill in resume.skills:
        by_category.setdefault(skill.category or "General", []).append(skill.name)
    return by_category


def render_resume_pdf(resume: Resume) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch
    )
    styles = getSampleStyleSheet()
    heading = styles["Heading2"]
    body = styles["BodyText"]
    name_style = ParagraphStyle("Name", parent=styles["Title"], alignment=0)

    full_name = resume.personal_info.full_name if resume.personal_info else resume.name
    story: list[object] = [Paragraph(full_name, name_style)]

    contact = _contact_line(resume)
    if contact:
        story.append(Paragraph(contact, body))
    if resume.summary:
        story.append(Spacer(1, 12))
        story.append(Paragraph(resume.summary, body))

    if resume.skills:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Skills", heading))
        for category, names in _skills_by_category(resume).items():
            story.append(Paragraph(f"<b>{category}:</b> {', '.join(names)}", body))

    if resume.experience:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Experience", heading))
        for exp in resume.experience:
            story.append(
                Paragraph(f"<b>{exp.role}</b>, {exp.company} ({_experience_dates(exp)})", body)
            )
            if exp.highlights:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(h, body)) for h in exp.highlights],
                        bulletType="bullet",
                    )
                )

    if resume.projects:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Projects", heading))
        for project in resume.projects:
            title = project.name
            if project.tech_stack:
                title += f" — {', '.join(project.tech_stack)}"
            story.append(Paragraph(f"<b>{title}</b>", body))
            if project.description:
                story.append(Paragraph(project.description, body))

    if resume.education:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Education", heading))
        for edu in resume.education:
            line = f"<b>{edu.degree}</b>, {edu.institution}"
            if edu.field_of_study:
                line += f" — {edu.field_of_study}"
            story.append(Paragraph(line, body))

    if resume.certifications:
        story.append(Spacer(1, 12))
        story.append(Paragraph("Certifications", heading))
        for cert in resume.certifications:
            line = cert.name
            if cert.issuer:
                line += f", {cert.issuer}"
            story.append(Paragraph(line, body))

    doc.build(story)
    return buffer.getvalue()


def render_resume_docx(resume: Resume) -> bytes:
    document = Document()
    full_name = resume.personal_info.full_name if resume.personal_info else resume.name
    document.add_heading(full_name, level=0)

    contact = _contact_line(resume)
    if contact:
        document.add_paragraph(contact)
    if resume.summary:
        document.add_paragraph(resume.summary)

    if resume.skills:
        document.add_heading("Skills", level=1)
        for category, names in _skills_by_category(resume).items():
            document.add_paragraph(f"{category}: {', '.join(names)}")

    if resume.experience:
        document.add_heading("Experience", level=1)
        for exp in resume.experience:
            document.add_paragraph(f"{exp.role}, {exp.company} ({_experience_dates(exp)})")
            for highlight in exp.highlights:
                document.add_paragraph(highlight, style="List Bullet")

    if resume.projects:
        document.add_heading("Projects", level=1)
        for project in resume.projects:
            title = project.name
            if project.tech_stack:
                title += f" — {', '.join(project.tech_stack)}"
            document.add_paragraph(title)
            if project.description:
                document.add_paragraph(project.description)

    if resume.education:
        document.add_heading("Education", level=1)
        for edu in resume.education:
            line = f"{edu.degree}, {edu.institution}"
            if edu.field_of_study:
                line += f" — {edu.field_of_study}"
            document.add_paragraph(line)

    if resume.certifications:
        document.add_heading("Certifications", level=1)
        for cert in resume.certifications:
            line = cert.name
            if cert.issuer:
                line += f", {cert.issuer}"
            document.add_paragraph(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
